from pathlib import Path

import pytest
from openpyxl import load_workbook

from evm_inventory.models import ConfigError
from evm_inventory.workbook import (
    WORKSHEET_HEADERS,
    WORKSHEET_NAME,
    WalletWorkbookRow,
    create_wallet_template,
    dry_run_workbook,
    load_wallet_workbook,
    parse_ordinal_ranges,
    wallet_range_batches,
)


def test_create_wallet_template_has_required_columns(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"

    create_wallet_template(output)

    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    assert tuple(cell.value for cell in sheet[1]) == WORKSHEET_HEADERS


def test_dry_run_writes_actions_without_changing_private_key(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    secret = "0x" + "a" * 64
    sheet.append([1, "0x" + "b" * 40, secret, "0x" + "c" * 40])
    workbook.save(output)

    summary = dry_run_workbook(output)

    assert summary == {"wallets": 1, "actions_written": 1}
    updated = load_workbook(output)[WORKSHEET_NAME]
    assert updated.cell(row=2, column=3).value == secret
    assert updated.cell(row=2, column=5).value == "Маршруты еще не рассчитаны"


def test_dry_run_rejects_invalid_deposit_address_without_writing_actions(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    sheet.append([1, "0x" + "b" * 40, "0x" + "a" * 64, "not-an-address", "keep"])
    workbook.save(output)

    with pytest.raises(ConfigError, match="invalid Bitget EVM deposit address"):
        dry_run_workbook(output)

    assert load_workbook(output)[WORKSHEET_NAME].cell(row=2, column=5).value == "keep"


def test_read_only_planning_accepts_blank_deposit_address(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    workbook[WORKSHEET_NAME].append([1, "0x" + "b" * 40, "0x" + "a" * 64, ""])
    workbook.save(output)

    rows = load_wallet_workbook(output, require_deposit_address=False)

    assert rows[0].bitget_deposit_address == ""


def test_rabby_proxy_is_required_per_wallet_and_never_shown_in_repr(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    proxy = "http://login:secret@proxy.example:8080"
    sheet.append([1, "0x" + "b" * 40, "0x" + "a" * 64, "", "", proxy])
    sheet.append([2, "0x" + "c" * 40, "0x" + "d" * 64, "", "", ""])
    workbook.save(output)

    rows = load_wallet_workbook(output, require_deposit_address=False)
    assert rows[0].rabby_proxy == proxy
    assert proxy not in repr(rows[0])
    with pytest.raises(ConfigError, match="row 3: Rabby proxy is required"):
        load_wallet_workbook(output, require_deposit_address=False, require_rabby_proxy=True)


def test_rabby_proxy_rejects_invalid_url_without_echoing_credentials(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    workbook[WORKSHEET_NAME].append(
        [1, "0x" + "b" * 40, "0x" + "a" * 64, "", "", "http://login:secret@proxy.example"]
    )
    workbook.save(output)

    with pytest.raises(ConfigError, match="row 2: invalid Rabby proxy") as error:
        load_wallet_workbook(output, require_deposit_address=False, require_rabby_proxy=True)
    assert "secret" not in str(error.value)


def test_rabby_proxy_accepts_socks5_url_and_rejects_scheme_less_value(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    sheet.append([1, "0x" + "b" * 40, "0x" + "a" * 64, "", "", "socks5://user:pass@proxy.example:1080"])
    sheet.append([2, "0x" + "c" * 40, "0x" + "d" * 64, "", "", "proxy.example:8080:user:pass"])
    workbook.save(output)

    with pytest.raises(ConfigError, match="row 3: invalid Rabby proxy"):
        load_wallet_workbook(output, require_deposit_address=False, require_rabby_proxy=True)
    book = load_workbook(output)
    book[WORKSHEET_NAME].delete_rows(3)
    book.save(output)
    rows = load_wallet_workbook(output, require_deposit_address=False, require_rabby_proxy=True)
    assert rows[0].rabby_proxy.startswith("socks5://")


def test_parse_ranges_and_group_rows():
    assert parse_ordinal_ranges("1-2, 5, 8-9") == ((1, 2), (5, 5), (8, 9))
    rows = tuple(
        WalletWorkbookRow(index + 2, ordinal, f"0x{ordinal:040x}", "0x" + "a" * 64, "")
        for index, ordinal in enumerate((1, 2, 5, 8, 9))
    )
    batches = wallet_range_batches(rows, parse_ordinal_ranges("1-2,8-9"))
    assert [[row.ordinal for row in batch] for batch in batches] == [[1, 2], [8, 9]]


def test_selected_workbook_range_ignores_invalid_unselected_rows(tmp_path: Path):
    output = tmp_path / "wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    workbook[WORKSHEET_NAME].append([1, "0x" + "b" * 40, "0x" + "a" * 64, "0x" + "c" * 40])
    workbook[WORKSHEET_NAME].append([2, "not-an-address", "bad-key", ""])
    workbook.save(output)

    rows = load_wallet_workbook(
        output,
        ordinal_ranges=parse_ordinal_ranges("1"),
    )

    assert [row.ordinal for row in rows] == [1]


def _proxy_workbook(tmp_path, ordinals=(1, 2, 4, 9), reserve_count=6):
    output = tmp_path / "proxy-wallets.xlsx"
    create_wallet_template(output)
    workbook = load_workbook(output)
    sheet = workbook[WORKSHEET_NAME]
    for index, ordinal in enumerate(ordinals):
        sheet.append([ordinal, f"0x{index + 1:040x}", "0x" + "a" * 64,
                      f"0x{index + 11:040x}", "", f"http://user:secret@primary{index}:8000"])
        sheet.append([None] * 6)
    for index in range(reserve_count):
        sheet.append([None, None, None, None, None, f"http://user:secret@reserve{index}:8000"])
    workbook.save(output)
    workbook.close()
    return output


def test_proxy_assignments_follow_physical_wallet_order_with_gaps(tmp_path):
    from evm_inventory.workbook import load_proxy_assignments

    path = _proxy_workbook(tmp_path)
    assignments = load_proxy_assignments(path)
    wallet4 = f"0x{3:040x}"
    assert assignments[wallet4] == (
        "http://user:secret@primary2:8000", "http://user:secret@reserve2:8000"
    )
    assert assignments[f"0x{1:040x}"][-1] == "http://user:secret@reserve4:8000"
    assert assignments[f"0x{2:040x}"][-1] == "http://user:secret@reserve5:8000"
    selected = load_wallet_workbook(path, ordinal_ranges=((4, 4),), require_rabby_proxy=True)
    assert len(selected) == 1
    assert selected[0].proxy_pool == assignments[wallet4]
    assert assignments[selected[0].public_address] == assignments[wallet4]
    assert len(load_wallet_workbook(path, require_rabby_proxy=True)) == 4


def test_legacy_proxy_assignments_require_no_reserves(tmp_path):
    from evm_inventory.workbook import load_proxy_assignments

    path = _proxy_workbook(tmp_path, reserve_count=0)
    assert list(map(len, load_proxy_assignments(path).values())) == [1] * 4


def test_proxy_assignments_do_not_validate_private_keys(tmp_path):
    from evm_inventory.workbook import load_proxy_assignments

    path = _proxy_workbook(tmp_path)
    book = load_workbook(path)
    book[WORKSHEET_NAME].cell(2, 3, "unselected-secret-invalid-key")
    book.save(path)
    assert len(load_proxy_assignments(path)) == 4
    assert len(load_wallet_workbook(path, ordinal_ranges=((4, 4),))) == 1


@pytest.mark.parametrize("change,expected", [
    ("incomplete", "reserve"),
    ("duplicate", "duplicate"),
    ("between", "proxy-only"),
    ("before", "proxy-only"),
    ("partial", "row 10"),
    ("multiline", "invalid Rabby proxy"),
    ("identity", "invalid public EVM address"),
    ("ordinal", "duplicate #"),
])
def test_proxy_assignment_errors_are_sanitized(tmp_path, change, expected):
    from evm_inventory.workbook import load_proxy_assignments

    path = _proxy_workbook(tmp_path)
    book = load_workbook(path)
    sheet = book[WORKSHEET_NAME]
    if change == "incomplete":
        sheet.delete_rows(12, 4)
    elif change == "duplicate":
        sheet.cell(10, 6, sheet.cell(2, 6).value)
    elif change == "between":
        sheet.cell(3, 6, "http://secret@between:8000")
    elif change == "before":
        sheet.insert_rows(2)
        sheet.cell(2, 6, "http://secret@before:8000")
    elif change == "partial":
        sheet.cell(10, 3, "private-secret")
    elif change == "multiline":
        sheet.cell(10, 6, "http://secret@reserve:8000\nhttp://secret@second:8000")
    elif change == "identity":
        sheet.cell(2, 2, "private-secret")
    elif change == "ordinal":
        sheet.cell(4, 1, 1)
    book.save(path)
    with pytest.raises(ConfigError, match=expected) as error:
        load_proxy_assignments(path)
    assert "secret" not in str(error.value)
    assert "http://" not in str(error.value)


def test_dry_run_validates_reserve_pool_before_writing(tmp_path):
    path = _proxy_workbook(tmp_path, reserve_count=1)
    with pytest.raises(ConfigError, match="reserve"):
        dry_run_workbook(path)
    assert load_workbook(path)[WORKSHEET_NAME].cell(2, 5).value is None


def test_dry_run_preserves_valid_reserve_rows(tmp_path):
    path = _proxy_workbook(tmp_path)
    assert dry_run_workbook(path)["wallets"] == 4
    assert load_workbook(path)[WORKSHEET_NAME].cell(10, 5).value is None


def test_required_legacy_proxy_pool_rejects_duplicate_primaries(tmp_path):
    path = _proxy_workbook(tmp_path, reserve_count=0)
    book = load_workbook(path)
    sheet = book[WORKSHEET_NAME]
    sheet.cell(4, 6, sheet.cell(2, 6).value)
    book.save(path)
    with pytest.raises(ConfigError, match="row 4: duplicate Rabby proxy"):
        load_wallet_workbook(path, require_rabby_proxy=True, ordinal_ranges=((4, 4),))


def test_proxy_assignment_helper_closes_workbook_on_validation_error(tmp_path, monkeypatch):
    import evm_inventory.workbook as module

    path = _proxy_workbook(tmp_path, reserve_count=1)
    book = load_workbook(path, read_only=True, data_only=True)
    original_close = book.close
    closed = []

    def close():
        closed.append(True)
        original_close()

    monkeypatch.setattr(book, "close", close)
    monkeypatch.setattr(module, "load_workbook", lambda *args, **kwargs: book)
    with pytest.raises(ConfigError, match="incomplete first reserve set"):
        module.load_proxy_assignments(path)
    assert closed == [True]
