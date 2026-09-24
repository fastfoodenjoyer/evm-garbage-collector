import hashlib
import json

from openpyxl import load_workbook

from evm_inventory.cli import main
from evm_inventory.store import Store
from evm_inventory.workbook import create_wallet_template

WALLET = "0x" + "1" * 40
ROUTER = "0x" + "2" * 40
TOKEN = "0x" + "3" * 40
PROXY = "http://login:secret@proxy.example:8080"


def workbook_with_proxy(path, *wallets):
    create_wallet_template(path)
    book = load_workbook(path)
    for ordinal, wallet in enumerate(wallets, start=1):
        book["Wallets"].append([ordinal, wallet, "0x" + "a" * 64, "", "", PROXY])
    book.save(path)


class FakeRabby:
    def __init__(self, *_args, **_kwargs):
        self.client = _args[0]

    def chain_ids(self):
        return {"eth": 1}

    def positions(self, _wallet):
        return [
            {
                "id": "vault",
                "chain": "eth",
                "portfolio_item_list": [
                    {
                        "name": "Deposit",
                        "pool": {"id": "pool", "chain": "eth"},
                        "position_index": "",
                        "proxy_detail": {},
                        "stats": {"net_usd_value": 12, "debt_usd_value": 0},
                        "detail": {"supply_token_list": [{"id": TOKEN, "amount": 12}]},
                        "withdraw_actions": [
                            {
                                "type": "withdraw",
                                "contract_id": ROUTER,
                                "func": "withdraw(uint256,address)",
                                "str_params": ["12", WALLET],
                                "need_approve": {},
                            }
                        ],
                    }
                ],
            }
        ]


def test_quote_defi_writes_key_free_plan_and_digest(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("evm_inventory.defi_operations.RabbyClient", FakeRabby)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n")
    output = tmp_path / "out" / "defi.json"
    db = tmp_path / "inventory.sqlite"
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET)
    assert main([
        "quote-defi", "--wallets", str(wallets), "--workbook", str(workbook_path),
        "--output", str(output), "--db", str(db)
    ]) == 0
    result = json.loads(capsys.readouterr().out)
    plan = json.loads(output.read_text())
    assert plan["summary"]["ready"] == 1
    assert result["plan_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert "private_key" not in output.read_text()
    assert "secret" not in output.read_text()
    with Store(db, readonly=True) as store:
        assert store.defi_plan(result["plan_sha256"]) == output.read_bytes()
        assert store.defi_positions(result["plan_sha256"])[0]["status"] == "ready"


def test_execute_defi_requires_explicit_action_and_reviewed_digest(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("evm_inventory.defi_operations.RabbyClient", FakeRabby)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n")
    plan_path = tmp_path / "defi.json"
    db = tmp_path / "inventory.sqlite"
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET)
    assert main([
        "quote-defi", "--wallets", str(wallets), "--output", str(plan_path),
        "--db", str(db), "--workbook", str(workbook_path),
    ]) == 0
    capsys.readouterr()
    plan = json.loads(plan_path.read_text())
    action_id = plan["entries"][0]["action_id"]
    called = {}

    def fake_execute(plan_bytes, **kwargs):
        called.update(kwargs)
        return {"status": "preview", "action_id": kwargs["action_id"]}

    monkeypatch.setattr("evm_inventory.defi_operations.execute_defi_action", fake_execute)
    seen_proxies = []

    class FakeClient:
        def __init__(self, **kwargs):
            if "proxy" in kwargs:
                seen_proxies.append((kwargs.get("proxy"), kwargs.get("trust_env")))

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            pass

        def close(self):
            pass

    monkeypatch.setattr("evm_inventory.defi_operations.httpx.Client", FakeClient)
    digest = hashlib.sha256(plan_path.read_bytes()).hexdigest()
    assert (
        main(
            [
                "execute-defi",
                "--plan",
                str(plan_path),
                "--plan-sha256",
                digest,
                "--action-id",
                action_id,
                "--workbook",
                str(workbook_path),
                "--db",
                str(db),
                "--max-gas-wei",
                "1000000",
            ]
        )
        == 0
    )
    assert called["execute"] is False
    assert called["action_id"] == action_id
    assert called["journal_path"] == db
    assert seen_proxies == [(PROXY, False)]
    assert json.loads(capsys.readouterr().out)["status"] == "preview"
    with Store(db, readonly=True) as store:
        assert store.defi_execution_events(digest)[0]["result"]["status"] == "preview"


def test_quote_defi_uses_matching_proxy_for_each_wallet(tmp_path, monkeypatch, capsys):
    second_wallet = "0x" + "4" * 40
    urls = {
        WALLET: "http://one:secret-one@proxy-one.example:8001",
        second_wallet: "http://two:secret-two@proxy-two.example:8002",
    }
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET, second_wallet)
    book = load_workbook(workbook_path)
    book["Wallets"].cell(row=2, column=6, value=urls[WALLET])
    book["Wallets"].cell(row=3, column=6, value=urls[second_wallet])
    book.save(workbook_path)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n" + second_wallet + "\n")
    used = []

    class FakeClient:
        def __init__(self, **kwargs):
            assert kwargs["trust_env"] is False
            self.proxy = kwargs["proxy"]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            pass

    class RoutedRabby:
        def __init__(self, client):
            self.proxy = client.proxy

        def chain_ids(self):
            return {"eth": 1}

        def positions(self, wallet):
            used.append((wallet, self.proxy))
            return []

    monkeypatch.setattr("evm_inventory.defi_operations.httpx.Client", FakeClient)
    monkeypatch.setattr("evm_inventory.defi_operations.RabbyClient", RoutedRabby)
    assert main([
        "quote-defi", "--wallets", str(wallets), "--workbook", str(workbook_path),
        "--output", str(tmp_path / "plan.json"), "--db", str(tmp_path / "inventory.sqlite"),
    ]) == 0
    capsys.readouterr()
    assert used == list(urls.items())


def test_quote_defi_accepts_socks5_proxy(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("evm_inventory.defi_operations.RabbyClient", FakeRabby)
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET)
    book = load_workbook(workbook_path)
    book["Wallets"].cell(row=2, column=6, value="socks5://user:pass@proxy.example:1080")
    book.save(workbook_path)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n")

    assert main([
        "quote-defi", "--wallets", str(wallets), "--workbook", str(workbook_path),
        "--output", str(tmp_path / "plan.json"), "--db", str(tmp_path / "inventory.sqlite"),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "planned"


def test_quote_defi_rejects_missing_proxy_before_network(tmp_path, monkeypatch, capsys):
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET)
    book = load_workbook(workbook_path)
    book["Wallets"].cell(row=2, column=6).value = None
    book.save(workbook_path)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n")

    def unexpected_client(**_kwargs):
        raise AssertionError("Rabby request must not be sent directly")

    monkeypatch.setattr("evm_inventory.defi_operations.httpx.Client", unexpected_client)
    assert main([
        "quote-defi", "--wallets", str(wallets), "--workbook", str(workbook_path),
        "--output", str(tmp_path / "plan.json"), "--db", str(tmp_path / "inventory.sqlite"),
    ]) == 2
    assert "Rabby proxy is required" in capsys.readouterr().err


def test_execute_defi_rejects_missing_proxy_before_rpc(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("evm_inventory.defi_operations.RabbyClient", FakeRabby)
    workbook_path = tmp_path / "wallets.xlsx"
    workbook_with_proxy(workbook_path, WALLET)
    wallets = tmp_path / "wallets.txt"
    wallets.write_text(WALLET + "\n")
    plan_path = tmp_path / "plan.json"
    db = tmp_path / "inventory.sqlite"
    assert main([
        "quote-defi", "--wallets", str(wallets), "--workbook", str(workbook_path),
        "--output", str(plan_path), "--db", str(db),
    ]) == 0
    capsys.readouterr()
    action_id = json.loads(plan_path.read_text())["entries"][0]["action_id"]
    book = load_workbook(workbook_path)
    book["Wallets"].cell(row=2, column=6).value = None
    book.save(workbook_path)

    def unexpected_execute(*_args, **_kwargs):
        raise AssertionError("execution must stop before RPC")

    monkeypatch.setattr("evm_inventory.defi_operations.execute_defi_action", unexpected_execute)
    assert main([
        "execute-defi", "--plan", str(plan_path),
        "--plan-sha256", hashlib.sha256(plan_path.read_bytes()).hexdigest(),
        "--action-id", action_id, "--workbook", str(workbook_path),
        "--db", str(db), "--max-gas-wei", "1000000",
    ]) == 2
    assert "Rabby proxy is required" in capsys.readouterr().err
