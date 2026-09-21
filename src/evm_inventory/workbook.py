"""XLSX input and dry-run output for wallet operations."""

from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import ConfigError

WORKSHEET_NAME = "Wallets"
WORKSHEET_HEADERS = (
    "# (номер)",
    "Публичный адрес",
    "Приватный ключ",
    "Адрес депозита Bitget",
    "Предлагаемые действия",
)
_ADDRESS_RE = re.compile(r"^0x[0-9a-fA-F]{40}$")
_PRIVATE_KEY_RE = re.compile(r"^0x[0-9a-fA-F]{64}$")
_ORDINAL_RANGE_RE = re.compile(r"^(\d+)(?:-(\d+))?$")


@dataclass(frozen=True, slots=True)
class WalletWorkbookRow:
    row_number: int
    ordinal: int
    public_address: str
    private_key: str = field(repr=False)
    bitget_deposit_address: str


def parse_ordinal_ranges(spec: str | None) -> tuple[tuple[int, int], ...] | None:
    """Parse inclusive ordinal ranges such as ``1-50,75,100-120``."""

    if spec is None:
        return None
    parts = [part.strip() for part in spec.split(",")]
    if not parts or any(not part for part in parts):
        raise ConfigError("wallet ranges must be comma-separated positive numbers or ranges")
    ranges: list[tuple[int, int]] = []
    for part in parts:
        match = _ORDINAL_RANGE_RE.fullmatch(part)
        if match is None:
            raise ConfigError(f"invalid wallet range {part!r}")
        start = int(match.group(1))
        end = int(match.group(2) or match.group(1))
        if start <= 0 or end <= 0 or end < start:
            raise ConfigError(f"invalid wallet range {part!r}")
        ranges.append((start, end))
    for index, current in enumerate(ranges):
        for previous in ranges[:index]:
            if current[0] <= previous[1] and previous[0] <= current[1]:
                raise ConfigError("wallet ranges must not overlap")
    return tuple(ranges)


def wallet_range_batches(
    rows: tuple[WalletWorkbookRow, ...],
    ranges: tuple[tuple[int, int], ...] | None,
) -> tuple[tuple[WalletWorkbookRow, ...], ...]:
    """Return selected rows grouped by requested ranges, preserving sheet order."""

    if ranges is None:
        return (rows,)
    batches: list[tuple[WalletWorkbookRow, ...]] = []
    found: set[int] = set()
    for start, end in ranges:
        batch = tuple(row for row in rows if start <= row.ordinal <= end)
        found.update(row.ordinal for row in batch)
        batches.append(batch)
    requested = {ordinal for start, end in ranges for ordinal in range(start, end + 1)}
    missing = sorted(requested - found)
    if missing:
        preview = ", ".join(str(item) for item in missing[:10])
        suffix = "..." if len(missing) > 10 else ""
        raise ConfigError(f"wallet ordinal(s) not found in workbook: {preview}{suffix}")
    return tuple(batches)


def create_wallet_template(path: Path) -> None:
    """Create an empty wallet workbook with the required input columns."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKSHEET_NAME
    for index, header in enumerate(WORKSHEET_HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    for column, width in zip(("A", "B", "C", "D", "E"), (8, 46, 72, 46, 72), strict=True):
        sheet.column_dimensions[column].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _required_string(value: object, *, row_number: int, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ConfigError(f"row {row_number}: {field_name} is required")
    return value.strip()


def load_wallet_workbook(
    path: Path,
    *,
    require_deposit_address: bool = True,
    ordinal_ranges: tuple[tuple[int, int], ...] | None = None,
) -> tuple[WalletWorkbookRow, ...]:
    """Read and validate operator input without persisting private keys."""

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except OSError as exc:
        raise ConfigError(f"cannot read workbook {path}: {exc}") from exc
    if WORKSHEET_NAME not in workbook.sheetnames:
        raise ConfigError(f"workbook must contain a {WORKSHEET_NAME!r} sheet")
    sheet = workbook[WORKSHEET_NAME]
    headers = tuple(cell.value for cell in sheet[1])
    if headers != WORKSHEET_HEADERS:
        raise ConfigError("workbook headers do not match the wallet template")

    rows: list[WalletWorkbookRow] = []
    seen_ordinals: set[int] = set()
    seen_wallets: set[str] = set()
    issues: list[str] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or value == "" for value in values):
            continue
        try:
            ordinal = values[0]
            if isinstance(ordinal, bool) or not isinstance(ordinal, int) or ordinal <= 0:
                raise ConfigError(f"row {row_number}: # must be a positive integer")
            if ordinal_ranges is not None and not any(
                start <= ordinal <= end for start, end in ordinal_ranges
            ):
                continue
            if ordinal in seen_ordinals:
                raise ConfigError(f"row {row_number}: duplicate # {ordinal}")
            public_address = _required_string(
                values[1], row_number=row_number, field_name="public address"
            ).lower()
            private_key = _required_string(
                values[2], row_number=row_number, field_name="private key"
            )
            deposit_address = (
                _required_string(
                    values[3], row_number=row_number, field_name="Bitget deposit address"
                ).lower()
                if require_deposit_address or values[3]
                else ""
            )
            if not _ADDRESS_RE.fullmatch(public_address):
                raise ConfigError(f"row {row_number}: invalid public EVM address")
            if not _PRIVATE_KEY_RE.fullmatch(private_key):
                raise ConfigError(
                    f"row {row_number}: private key must be 0x plus 64 hex characters"
                )
            if deposit_address and not _ADDRESS_RE.fullmatch(deposit_address):
                raise ConfigError(f"row {row_number}: invalid Bitget EVM deposit address")
            if public_address in seen_wallets:
                raise ConfigError(f"row {row_number}: duplicate public address")
            seen_ordinals.add(ordinal)
            seen_wallets.add(public_address)
            rows.append(
                WalletWorkbookRow(
                    row_number=row_number,
                    ordinal=ordinal,
                    public_address=public_address,
                    private_key=private_key,
                    bitget_deposit_address=deposit_address,
                )
            )
        except ConfigError as exc:
            issues.extend(exc.issues)
    if issues:
        raise ConfigError("invalid wallet workbook: " + "; ".join(issues), issues=tuple(issues))
    if not rows:
        raise ConfigError("workbook must contain at least one wallet")
    return tuple(rows)


def dry_run_workbook(path: Path) -> dict[str, int]:
    """Validate all inputs, then atomically write placeholder plans to the workbook."""

    rows = load_wallet_workbook(path)
    workbook = load_workbook(path)
    sheet = workbook[WORKSHEET_NAME]
    for row in rows:
        sheet.cell(row=row.row_number, column=5, value="Маршруты еще не рассчитаны")
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return {"wallets": len(rows), "actions_written": len(rows)}
